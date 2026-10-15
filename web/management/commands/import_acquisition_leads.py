"""Importiert kurierte Akquise-Leads (Kleinbetriebe) aus einer Excel-Liste.

Liest das Blatt 'Leads' der Akquise-Liste und legt/aktualisiert
AcquisitionLead-Datensätze an (Dedupe über Firma+Ort). Für Datalization.

    python manage.py import_acquisition_leads "/pfad/Akquise_Zofingen_Kleinbetriebe.xlsx"
    python manage.py import_acquisition_leads <datei> --sheet Leads --dry-run
"""
import re

from django.core.management.base import BaseCommand, CommandError

from web.models import AcquisitionLead


def _s(v):
    return ("" if v is None else str(v)).strip()


class Command(BaseCommand):
    help = "Importiert Akquise-Leads (Kleinbetriebe) aus einer .xlsx-Liste."

    def add_arguments(self, parser):
        parser.add_argument("path", help="Pfad zur .xlsx-Datei")
        parser.add_argument("--sheet", default="Leads", help="Blattname (Default: Leads)")
        parser.add_argument("--dry-run", action="store_true", help="Nur zählen, nichts speichern.")

    def handle(self, *args, **opts):
        try:
            import openpyxl
        except ImportError:
            raise CommandError("openpyxl fehlt – bitte `pip install openpyxl`.")

        wb = openpyxl.load_workbook(opts["path"], read_only=True, data_only=True)
        if opts["sheet"] not in wb.sheetnames:
            raise CommandError(f"Blatt '{opts['sheet']}' nicht gefunden. "
                               f"Vorhanden: {wb.sheetnames}")
        ws = wb[opts["sheet"]]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise CommandError("Blatt ist leer.")
        hdr = [_s(h) for h in rows[0]]
        idx = {h: i for i, h in enumerate(hdr)}

        def col(row, name):
            i = idx.get(name)
            return _s(row[i]) if i is not None and i < len(row) else ""

        data = [r for r in rows[1:] if r and col(r, "Firma")]
        self.stdout.write(f"Datenzeilen: {len(data)}")

        created = updated = 0
        for r in data:
            name = col(r, "Firma")
            city = col(r, "Ort")
            key = re.sub(r"[^a-z0-9]", "", name.lower()) + "|" + city.lower()
            defaults = {
                "name": name[:255],
                "segment": col(r, "Segment")[:160],
                "city": city[:120],
                "address": col(r, "Adresse")[:300],
                "phone": col(r, "Telefon")[:60],
                "priority": col(r, "Prio")[:20],
                "size_note": col(r, "Grösse")[:200],
                "employees": col(r, "Mitarbeitende (Schätzung)")[:80],
                "size_source": col(r, "Quelle Grösse")[:200],
                "approach": col(r, "Ansatz (Gesprächseinstieg)"),
            }
            if opts["dry_run"]:
                continue
            _, is_new = AcquisitionLead.objects.update_or_create(ext_key=key, defaults=defaults)
            created += is_new
            updated += (not is_new)

        if opts["dry_run"]:
            self.stdout.write(self.style.WARNING(
                f"DRY-RUN: {len(data)} Leads erkannt. Nichts gespeichert."))
        else:
            self.stdout.write(self.style.SUCCESS(
                f"Importiert: {created} neu, {updated} aktualisiert. "
                f"Total: {AcquisitionLead.objects.count()}"))
